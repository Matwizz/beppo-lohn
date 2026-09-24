"""
beppo AG – Payroll Cost Center Tool
Flask backend: accepts Bexio lohnkonto export, returns booking Excel.
"""
from flask import Flask, request, send_file, render_template
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import os
import tempfile

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max upload

# ── KST mapping (hardcoded – update when org changes) ────────────────────────
KST_DATA = [
    ('Stefan',      'Bjelajac',         400, 'IT'),
    ('Tobias',      'Boschung',         400, 'IT'),
    ('Justin',      'Jaeger',           400, 'IT'),
    ('Anina',       'Haller',           300, 'Services'),
    ('Juan',        'Huber',            300, 'Services'),
    ('Benoit',      'Corot',            200, 'Sales'),
    ('Madeleine',   'Ricklin-Kuschel',  200, 'Sales'),
    ('Yakup',       'Tasdemir',         200, 'Sales'),
    ('Isaia',       'Valeri',           200, 'Sales'),
    ('Igor',        'Vidackovic',       200, 'Sales'),
    ('Christopher', 'Horak',            200, 'Sales'),
    ('Marco',       'Muscarello',       100, 'Management'),
    ('Philip',      'Ricklin',          100, 'Management'),
    ('David',       'Shulman',          100, 'Management'),
    ('Gökhan',      'Tüzün',            100, 'Management'),
]

KST_BY_LAST = {last.lower(): (did, dname) for (_, last, did, dname) in KST_DATA}

DEPT_ORDER = [(100,'Management'),(200,'Sales'),(300,'Services'),(400,'IT')]

KST_FILLS = {
    100: PatternFill("solid", fgColor="E2EFDA"),
    200: PatternFill("solid", fgColor="DDEBF7"),
    300: PatternFill("solid", fgColor="FFF2CC"),
    400: PatternFill("solid", fgColor="FCE4D6"),
}

BOOKING_LINES = [
    (1,  '1091','2271','AHV/IV/EO & ALV Beitrag AN',               (5010,5020)),
    (2,  '1091','2270','BVG Beitrag AN',                            (5050,)),
    (3,  '1091','2273','NBU Beitrag AN',                            (5040,)),
    (4,  '1091','2274','KTG Beitrag AN',                            (5045,)),
    (5,  '1091','1090','Auszahlungsbetrag',                         (6600,)),
    (6,  '5000','1091','Monatslohn / Bruttolohn',                   (5000,)),
    (7,  '5003','1091','Provision',                                  (1218,)),
    (8,  '5700','2271','AHV VK, ALV AG, AHV/IV/EO AG',             (7010,7011,7020)),
    (9,  '5710','2271','FAK Beitrag AG',                            (7070,)),
    (10, '5720','2270','BVG Beitrag AG',                            (7050,)),
    (11, '5730','2273','BU AG, UVGZ AG, NBU AG',                   (7040,7041)),
    (12, '5740','2274','KTG Beitrag AG',                            (7045,)),
    (13, '5822','1091','Effektive Spesen',                          (6000,)),
    (14, '5891','1091','Ausgleich Geschäftswagen',                  (5100,)),
    (15, '1091','6270','Privatanteil Geschäftswagen (inkl. MwSt.)',(1910,)),
]

RELEVANT_LA = {la for line in BOOKING_LINES for la in line[4]}

MONTH_COLS = {'Juli': 9, 'August': 10, 'September': 11}

HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
HDR_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))
CHF_FMT = '#,##0.00'

def find_dept(sheet_name):
    words = sheet_name.lower().split()
    for word in reversed(words):
        if word in KST_BY_LAST:
            return KST_BY_LAST[word]
    return None  # excluded (e.g. Sulamithe)

def build_excel(lohn_file_bytes, year=2026):
    lohn_wb = load_workbook(io.BytesIO(lohn_file_bytes), read_only=True, data_only=True)

    dept_totals = {m: {d[0]: {} for d in DEPT_ORDER} for m in MONTH_COLS}

    for sheet_name in lohn_wb.sheetnames:
        dept = find_dept(sheet_name)
        if dept is None:
            continue
        did = dept[0]
        ws   = lohn_wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows) if r and r[1] == 'LA'), None)
        if header_idx is None:
            continue
        for row in rows[header_idx + 1:]:
            if not row or row[1] is None:
                continue
            try:
                la = int(row[1])
            except (ValueError, TypeError):
                continue
            if la not in RELEVANT_LA:
                continue
            for month, col in MONTH_COLS.items():
                try:
                    val = float(row[col]) if row[col] not in (None, '') else 0.0
                except (ValueError, TypeError):
                    val = 0.0
                if abs(val) > 0:
                    dept_totals[month][did][la] = dept_totals[month][did].get(la, 0.0) + val

    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    # Instructions tab
    ws_i = wb_out.create_sheet("Instructions", 0)
    ws_i.column_dimensions['A'].width = 80
    lines = [
        (f"beppo AG – Payroll Booking File with Cost Centers ({year})", True, 13),
        ("", False, 10),
        ("HOW TO USE THIS TOOL EACH MONTH", True, 11),
        ("1. Export the payroll file from Bexio: all employees, same format as 'lohnkonto_alle_mitarbeiter'.", False, 10),
        ("2. Go to the Mathias.AI tool (URL provided by Trust Work).", False, 10),
        ("3. Upload the Bexio export and click 'Generate Booking File'.", False, 10),
        ("4. Download the Excel and copy the entries into Business Central.", False, 10),
        ("", False, 10),
        ("COST CENTER MAPPING", True, 11),
        ("KST 100 – Management:  Marco Muscarello, Philip Ricklin, David Shulman, Gökhan Tüzün", False, 10),
        ("KST 200 – Sales:       Benoit Corot, Madeleine Ricklin-Kuschel, Yakup Tasdemir, Isaia Valeri, Igor Vidackovic, Christopher Horak", False, 10),
        ("KST 300 – Services:    Anina Haller, Juan Bautista Huber", False, 10),
        ("KST 400 – IT:          Stefan Bjelajac, Tobias Boschung, Justin Jaeger", False, 10),
        ("Note: Sulamithe Perrenoud is excluded (left beppo as of August 2026).", False, 10),
        ("", False, 10),
        ("NOTES", True, 11),
        ("- Zero-amount lines are hidden to keep the file clean.", False, 10),
        ("- September 2026: payroll run was not yet closed at time of first generation.", False, 10),
        ("- Generated by Mathias.AI | Trust Work AG | Switzerland", False, 9),
    ]
    for i, (text, bold, size) in enumerate(lines, 1):
        ws_i.cell(row=i, column=1, value=text).font = Font(name="Arial", bold=bold, size=size)
        ws_i.row_dimensions[i].height = 16 if text else 5

    for month in MONTH_COLS:
        ws = wb_out.create_sheet(f"{month} {year}")
        ws.merge_cells('A1:G1')
        ws['A1'] = f"beppo AG – Payroll Booking by Cost Center | {month} {year}"
        ws['A1'].font = Font(name="Arial", bold=True, size=12)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 22
        ws.merge_cells('A2:G2')
        ws['A2'] = "Generated by Mathias.AI | Trust Work AG – Switzerland"
        ws['A2'].font = Font(name="Arial", italic=True, size=9, color="888888")
        ws.row_dimensions[3].height = 6

        headers = ['No','Cost Center','Debit','Credit','Description','Amount (CHF)','Dept.']
        widths  = [5, 14, 8, 8, 48, 16, 14]
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = THIN_BORDER
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[4].height = 18

        row_num = 5
        grand_total = 0.0

        for did, dname in DEPT_ORDER:
            la_data = dept_totals[month].get(did, {})
            fill = KST_FILLS[did]

            ws.merge_cells(f'A{row_num}:G{row_num}')
            ws[f'A{row_num}'] = f"  KST {did} – {dname}"
            ws[f'A{row_num}'].font = Font(name="Arial", bold=True, size=10, color="1F4E79")
            ws[f'A{row_num}'].fill = fill
            ws[f'A{row_num}'].alignment = Alignment(vertical='center')
            ws.row_dimensions[row_num].height = 16
            row_num += 1

            kst_total = 0.0
            for line_no, soll, haben, desc, la_codes in BOOKING_LINES:
                amount = abs(sum(la_data.get(la, 0.0) for la in la_codes))
                if amount == 0:
                    continue
                for c, val in enumerate([line_no, did, soll, haben, desc, round(amount,2), dname], 1):
                    cell = ws.cell(row=row_num, column=c, value=val)
                    cell.font  = BODY_FONT
                    cell.border= THIN_BORDER
                    cell.alignment = Alignment(vertical='center')
                ws.cell(row=row_num, column=1).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(row=row_num, column=6).number_format = CHF_FMT
                ws.cell(row=row_num, column=6).alignment = Alignment(horizontal='right', vertical='center')
                ws.row_dimensions[row_num].height = 15
                kst_total += amount
                row_num += 1

            # Subtotal row
            for c in range(1, 8):
                ws.cell(row=row_num, column=c).fill = fill
                ws.cell(row=row_num, column=c).border = THIN_BORDER
            ws.cell(row=row_num, column=5, value=f"Subtotal KST {did} – {dname}").font = BOLD_FONT
            ws.cell(row=row_num, column=5).fill = fill
            ws.cell(row=row_num, column=6, value=round(kst_total,2)).font = BOLD_FONT
            ws.cell(row=row_num, column=6).fill = fill
            ws.cell(row=row_num, column=6).number_format = CHF_FMT
            ws.cell(row=row_num, column=6).alignment = Alignment(horizontal='right', vertical='center')
            grand_total += kst_total
            row_num += 2

        # Grand total
        ws.cell(row=row_num, column=5, value="TOTAL").font = Font(name="Arial", bold=True, size=11)
        ws.cell(row=row_num, column=5).border = THIN_BORDER
        ws.cell(row=row_num, column=6, value=round(grand_total,2)).font = Font(name="Arial", bold=True, size=11)
        ws.cell(row=row_num, column=6).number_format = CHF_FMT
        ws.cell(row=row_num, column=6).alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(row=row_num, column=6).border = THIN_BORDER
        ws.row_dimensions[row_num].height = 18
        ws.freeze_panes = 'A5'

    output = io.BytesIO()
    wb_out.save(output)
    output.seek(0)
    return output

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate():
    if 'bexio_file' not in request.files:
        return {'error': 'No file uploaded'}, 400
    file = request.files['bexio_file']
    if not file.filename.endswith('.xlsx'):
        return {'error': 'Please upload an .xlsx file'}, 400
    try:
        excel_output = build_excel(file.read())
        return send_file(
            excel_output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='beppo_Payroll_CostCenters.xlsx'
        )
    except Exception as e:
        return {'error': str(e)}, 500

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
